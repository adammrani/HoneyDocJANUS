"""Génération d'un HoneyDoc Excel comptable crédible et sans macro.

Le LLM fournit la synthèse narrative. Les montants, écritures et formules sont
produits par du code déterministe afin de garantir l'équilibre comptable. Le
classeur ne contient ni VBA ni macro. Une URL Canarytoken peut être attachée
comme image externe OOXML afin de détecter l'ouverture.
"""

from __future__ import annotations

import hashlib
import random
from collections import defaultdict
from dataclasses import dataclass
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal, ROUND_HALF_UP

from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo

from src.janus.generators.xlsx_canary import (
    CanarytokenWorkbook,
    attach_placeholder_image,
)


GENERATOR_VERSION = "janus-accounting-xlsx/1.1.0"
_CURRENCY_FORMAT = '#,##0.00 [$€-fr-FR]'
_MONEY = Decimal("0.01")

_ACCOUNTS = {
    "101000": ("Capital social", "equity"),
    "401000": ("Fournisseurs", "liability"),
    "411000": ("Clients", "asset"),
    "445660": ("TVA déductible", "asset"),
    "445710": ("TVA collectée", "liability"),
    "512000": ("Banque", "asset"),
    "607000": ("Achats et prestations", "expense"),
    "613200": ("Locations immobilières", "expense"),
    "641000": ("Rémunérations du personnel", "expense"),
    "645000": ("Charges sociales", "expense"),
    "706000": ("Prestations de services", "revenue"),
}


@dataclass(frozen=True)
class JournalLine:
    """Une ligne d'écriture comptable en partie double."""

    posting_date: date
    document_no: str
    description: str
    account: str
    debit: Decimal = Decimal("0.00")
    credit: Decimal = Decimal("0.00")


def _money(value: Decimal | int | float | str) -> Decimal:
    return Decimal(str(value)).quantize(_MONEY, rounding=ROUND_HALF_UP)


def _append_transaction(
    lines: list[JournalLine],
    posting_date: date,
    document_no: str,
    description: str,
    postings: list[tuple[str, Decimal | int | float | str, str]],
) -> None:
    for account, amount, side in postings:
        value = _money(amount)
        lines.append(
            JournalLine(
                posting_date=posting_date,
                document_no=document_no,
                description=description,
                account=account,
                debit=value if side == "debit" else Decimal("0.00"),
                credit=value if side == "credit" else Decimal("0.00"),
            )
        )


def validate_balanced_journal(lines: list[JournalLine]) -> None:
    """Lève ``ValueError`` si une pièce ou le journal global est déséquilibré."""

    if not lines:
        raise ValueError("Le journal comptable est vide.")

    by_document: dict[str, list[Decimal]] = defaultdict(
        lambda: [Decimal("0.00"), Decimal("0.00")]
    )
    for line in lines:
        if line.account not in _ACCOUNTS:
            raise ValueError(f"Compte inconnu : {line.account}")
        if line.debit < 0 or line.credit < 0:
            raise ValueError("Un montant comptable ne peut pas être négatif.")
        if bool(line.debit) == bool(line.credit):
            raise ValueError(
                f"La ligne {line.document_no}/{line.account} doit avoir un seul côté."
            )
        by_document[line.document_no][0] += line.debit
        by_document[line.document_no][1] += line.credit

    for document_no, (debit, credit) in by_document.items():
        if debit.quantize(_MONEY) != credit.quantize(_MONEY):
            raise ValueError(
                f"Pièce {document_no} déséquilibrée : débit={debit}, crédit={credit}."
            )

    total_debit = sum((line.debit for line in lines), Decimal("0.00"))
    total_credit = sum((line.credit for line in lines), Decimal("0.00"))
    if total_debit.quantize(_MONEY) != total_credit.quantize(_MONEY):
        raise ValueError("Le journal global est déséquilibré.")


def _seed_from(token_id: str, fiscal_year: int) -> int:
    material = f"{token_id}|{fiscal_year}|JANUS-XLSX".encode("utf-8")
    return int.from_bytes(hashlib.sha256(material).digest()[:8], "big")


def _build_journal(fiscal_year: int, seed: int) -> list[JournalLine]:
    rng = random.Random(seed)
    lines: list[JournalLine] = []
    vat_rate = Decimal("0.20")

    _append_transaction(
        lines,
        date(fiscal_year, 1, 2),
        f"OUV-{fiscal_year}",
        "Apport initial des actionnaires",
        [("512000", 250_000, "debit"), ("101000", 250_000, "credit")],
    )

    for month in range(1, 13):
        base_date = date(fiscal_year, month, 5)
        suffix = f"{fiscal_year}{month:02d}"

        sales_net = _money(rng.randint(22_000, 38_000))
        sales_vat = _money(sales_net * vat_rate)
        sales_gross = sales_net + sales_vat
        _append_transaction(
            lines,
            base_date,
            f"FV-{suffix}-01",
            "Facturation prestations clients grands comptes",
            [
                ("411000", sales_gross, "debit"),
                ("706000", sales_net, "credit"),
                ("445710", sales_vat, "credit"),
            ],
        )
        _append_transaction(
            lines,
            base_date + timedelta(days=8),
            f"BQ-{suffix}-01",
            "Encaissement factures clients",
            [("512000", sales_gross, "debit"), ("411000", sales_gross, "credit")],
        )

        purchases_net = _money(rng.randint(7_500, 13_500))
        purchases_vat = _money(purchases_net * vat_rate)
        purchases_gross = purchases_net + purchases_vat
        _append_transaction(
            lines,
            base_date + timedelta(days=2),
            f"FA-{suffix}-01",
            "Achats et prestations sous-traitées",
            [
                ("607000", purchases_net, "debit"),
                ("445660", purchases_vat, "debit"),
                ("401000", purchases_gross, "credit"),
            ],
        )
        _append_transaction(
            lines,
            base_date + timedelta(days=13),
            f"BQ-{suffix}-02",
            "Règlement fournisseurs",
            [("401000", purchases_gross, "debit"), ("512000", purchases_gross, "credit")],
        )

        rent_net = Decimal("4500.00")
        rent_vat = _money(rent_net * vat_rate)
        rent_gross = rent_net + rent_vat
        _append_transaction(
            lines,
            base_date + timedelta(days=1),
            f"LO-{suffix}",
            "Loyer mensuel siège social",
            [
                ("613200", rent_net, "debit"),
                ("445660", rent_vat, "debit"),
                ("401000", rent_gross, "credit"),
            ],
        )
        _append_transaction(
            lines,
            base_date + timedelta(days=3),
            f"BQ-{suffix}-03",
            "Paiement du loyer",
            [("401000", rent_gross, "debit"), ("512000", rent_gross, "credit")],
        )

        salaries = _money(rng.randint(14_200, 16_800))
        social_charges = _money(salaries * Decimal("0.28"))
        payroll_date = date(fiscal_year, month, 25)
        _append_transaction(
            lines,
            payroll_date,
            f"PAIE-{suffix}",
            "Rémunérations mensuelles",
            [("641000", salaries, "debit"), ("512000", salaries, "credit")],
        )
        _append_transaction(
            lines,
            payroll_date + timedelta(days=1),
            f"SOC-{suffix}",
            "Charges sociales mensuelles",
            [
                ("645000", social_charges, "debit"),
                ("512000", social_charges, "credit"),
            ],
        )

    lines.sort(key=lambda item: (item.posting_date, item.document_no, item.account))
    validate_balanced_journal(lines)
    return lines


def _style_title(cell, fill: str = "17365D") -> None:
    cell.font = Font(name="Aptos Display", size=18, bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor=fill)
    cell.alignment = Alignment(vertical="center")


def _style_header(row) -> None:
    thin = Side(style="thin", color="A6A6A6")
    for cell in row:
        cell.font = Font(name="Aptos", bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(bottom=thin)


def _add_table(worksheet, reference: str, name: str) -> None:
    table = Table(displayName=name, ref=reference)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    worksheet.add_table(table)


def build_accounting_workbook(
    summary_text: str,
    token_id: str,
    token_url: str = "",
    company_name: str = "Atlas Conseil & Industrie SA",
    fiscal_year: int | None = None,
) -> Workbook:
    """Construit un classeur financier cohérent, macro-free et prêt à déployer."""

    year = fiscal_year or datetime.now().year
    lines = _build_journal(year, _seed_from(token_id, year))

    wb = CanarytokenWorkbook(token_url=token_url)
    ws_summary = wb.active
    ws_summary.title = "Synthèse"
    ws_journal = wb.create_sheet("Journal")
    ws_balance = wb.create_sheet("Balance")
    ws_income = wb.create_sheet("Résultat")
    ws_sheet = wb.create_sheet("Bilan")
    ws_budget = wb.create_sheet("Budget mensuel")
    ws_meta = wb.create_sheet("_Document_Metadata")

    wb.properties.creator = company_name
    wb.properties.title = f"Dossier financier confidentiel {year}"
    wb.properties.subject = "Reporting financier interne"
    wb.properties.keywords = "finance, reporting, confidentiel"
    wb.properties.description = "Dossier financier interne confidentiel."
    try:
        wb.calculation.fullCalcOnLoad = True
        wb.calculation.forceFullCalc = True
        wb.calculation.calcMode = "auto"
    except AttributeError:
        pass

    journal_headers = [
        "Date", "Pièce", "Libellé", "Compte", "Intitulé", "Débit", "Crédit"
    ]
    ws_journal.append(journal_headers)
    for line in lines:
        ws_journal.append(
            [
                line.posting_date,
                line.document_no,
                line.description,
                line.account,
                _ACCOUNTS[line.account][0],
                float(line.debit),
                float(line.credit),
            ]
        )
    _style_header(ws_journal[1])
    ws_journal.freeze_panes = "A2"
    # Le tableau possède déjà son propre AutoFilter. Ajouter aussi un filtre
    # au niveau de la feuille crée deux filtres sur la même plage ; certaines
    # versions d'Excel réparent alors le fichier en supprimant le tableau.
    _add_table(ws_journal, f"A1:G{ws_journal.max_row}", "JournalComptable")
    for cell in ws_journal["A"][1:]:
        cell.number_format = "dd/mm/yyyy"
    for column in ("F", "G"):
        for cell in ws_journal[column][1:]:
            cell.number_format = _CURRENCY_FORMAT
    for column, width in {
        "A": 13, "B": 17, "C": 44, "D": 13, "E": 28, "F": 16, "G": 16
    }.items():
        ws_journal.column_dimensions[column].width = width

    ws_balance.append(
        ["Compte", "Intitulé", "Total débit", "Total crédit", "Solde débiteur", "Solde créditeur"]
    )
    for row_index, account in enumerate(sorted(_ACCOUNTS), start=2):
        label = _ACCOUNTS[account][0]
        ws_balance.append(
            [
                account,
                label,
                f'=SUMIF(Journal!$D:$D,A{row_index},Journal!$F:$F)',
                f'=SUMIF(Journal!$D:$D,A{row_index},Journal!$G:$G)',
                f'=MAX(C{row_index}-D{row_index},0)',
                f'=MAX(D{row_index}-C{row_index},0)',
            ]
        )
    total_row = ws_balance.max_row + 1
    ws_balance.cell(total_row, 1, "TOTAL")
    for col in range(3, 7):
        letter = ws_balance.cell(1, col).column_letter
        ws_balance.cell(total_row, col, f"=SUM({letter}2:{letter}{total_row - 1})")
    _style_header(ws_balance[1])
    for cell in ws_balance[total_row]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAF7")
    for row in ws_balance.iter_rows(min_row=2, min_col=3, max_col=6):
        for cell in row:
            cell.number_format = _CURRENCY_FORMAT
    ws_balance.freeze_panes = "A2"
    ws_balance.column_dimensions["A"].width = 14
    ws_balance.column_dimensions["B"].width = 31
    for column in ("C", "D", "E", "F"):
        ws_balance.column_dimensions[column].width = 18

    balance_rows = {
        ws_balance.cell(row, 1).value: row for row in range(2, total_row)
    }

    ws_income.append(["COMPTE DE RÉSULTAT", None, None])
    ws_income.merge_cells("A1:C1")
    _style_title(ws_income["A1"], "548235")
    ws_income.row_dimensions[1].height = 30
    ws_income.append(["Poste", "Compte", "Montant"])
    _style_header(ws_income[2])
    income_items = [
        ("Chiffre d'affaires", "706000", "revenue"),
        ("Achats et prestations", "607000", "expense"),
        ("Locations", "613200", "expense"),
        ("Salaires", "641000", "expense"),
        ("Charges sociales", "645000", "expense"),
    ]
    for label, account, kind in income_items:
        balance_row = balance_rows[account]
        formula = (
            f"=Balance!F{balance_row}"
            if kind == "revenue"
            else f"=Balance!E{balance_row}"
        )
        ws_income.append([label, account, formula])
    ws_income.append(["Total produits", None, "=C3"])
    ws_income.append(["Total charges", None, "=SUM(C4:C7)"])
    ws_income.append(["RÉSULTAT NET", None, "=C8-C9"])
    for cell in ws_income[ws_income.max_row]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="548235")
    for cell in ws_income["C"][2:]:
        cell.number_format = _CURRENCY_FORMAT
    ws_income.column_dimensions["A"].width = 34
    ws_income.column_dimensions["B"].width = 15
    ws_income.column_dimensions["C"].width = 20

    ws_sheet.append(["BILAN SYNTHÉTIQUE", None, None])
    ws_sheet.merge_cells("A1:C1")
    _style_title(ws_sheet["A1"], "806000")
    ws_sheet.row_dimensions[1].height = 30
    ws_sheet.append(["Poste", "Compte", "Montant"])
    _style_header(ws_sheet[2])
    asset_accounts = ["512000", "411000", "445660"]
    liability_accounts = ["401000", "445710", "101000"]
    for account in asset_accounts:
        row = balance_rows[account]
        ws_sheet.append([_ACCOUNTS[account][0], account, f"=Balance!E{row}"])
    ws_sheet.append(["TOTAL ACTIF", None, "=SUM(C3:C5)"])
    liability_start = ws_sheet.max_row + 2
    ws_sheet.cell(liability_start, 1, "PASSIF ET CAPITAUX PROPRES")
    ws_sheet.cell(liability_start, 1).font = Font(bold=True, color="806000")
    for account in liability_accounts:
        row = balance_rows[account]
        ws_sheet.append([_ACCOUNTS[account][0], account, f"=Balance!F{row}"])
    ws_sheet.append(["Résultat net de l'exercice", None, "=Résultat!C10"])
    liability_total_row = ws_sheet.max_row + 1
    ws_sheet.append(
        [
            "TOTAL PASSIF ET CAPITAUX PROPRES",
            None,
            f"=SUM(C{liability_start + 1}:C{liability_total_row - 1})",
        ]
    )
    ws_sheet.append(["Contrôle actif - passif", None, f"=C6-C{liability_total_row}"])
    for cell in ws_sheet["C"][2:]:
        cell.number_format = _CURRENCY_FORMAT
    for row_number in (6, liability_total_row, liability_total_row + 1):
        for cell in ws_sheet[row_number]:
            cell.font = Font(bold=True)
    ws_sheet.column_dimensions["A"].width = 39
    ws_sheet.column_dimensions["B"].width = 15
    ws_sheet.column_dimensions["C"].width = 22

    ws_budget.append(["Mois", "CA réalisé", "Budget CA", "Charges réalisées", "Budget charges"])
    _style_header(ws_budget[1])
    month_names = [
        "Janvier", "Février", "Mars", "Avril", "Mai", "Juin",
        "Juillet", "Août", "Septembre", "Octobre", "Novembre", "Décembre",
    ]
    for month, month_name in enumerate(month_names, start=1):
        month_lines = [line for line in lines if line.posting_date.month == month]
        revenue = sum(
            (line.credit - line.debit for line in month_lines if line.account == "706000"),
            Decimal("0.00"),
        )
        expenses = sum(
            (
                line.debit - line.credit
                for line in month_lines
                if _ACCOUNTS[line.account][1] == "expense"
            ),
            Decimal("0.00"),
        )
        ws_budget.append(
            [
                month_name,
                float(revenue),
                float(_money(revenue * Decimal("1.04"))),
                float(expenses),
                float(_money(expenses * Decimal("0.98"))),
            ]
        )
    for row in ws_budget.iter_rows(min_row=2, min_col=2, max_col=5):
        for cell in row:
            cell.number_format = _CURRENCY_FORMAT
    ws_budget.freeze_panes = "A2"
    ws_budget.column_dimensions["A"].width = 16
    for column in ("B", "C", "D", "E"):
        ws_budget.column_dimensions[column].width = 20
    ws_budget.conditional_formatting.add(
        "B2:E13",
        ColorScaleRule(
            start_type="min", start_color="F8696B",
            mid_type="percentile", mid_value=50, mid_color="FFEB84",
            end_type="max", end_color="63BE7B",
        ),
    )
    chart = LineChart()
    chart.title = f"Réalisé et budget {year}"
    chart.y_axis.title = "Montant (€)"
    chart.x_axis.title = "Mois"
    chart.add_data(
        Reference(ws_budget, min_col=2, max_col=5, min_row=1, max_row=13),
        titles_from_data=True,
    )
    chart.set_categories(Reference(ws_budget, min_col=1, min_row=2, max_row=13))
    chart.height = 8
    chart.width = 17
    ws_budget.add_chart(chart, "G2")

    ws_summary.merge_cells("A1:F2")
    ws_summary["A1"] = f"DOSSIER FINANCIER CONFIDENTIEL — {company_name}"
    _style_title(ws_summary["A1"])
    ws_summary["A1"].alignment = Alignment(vertical="center", horizontal="left")
    ws_summary.row_dimensions[1].height = 28
    ws_summary["A4"] = "Exercice"
    ws_summary["B4"] = year
    ws_summary["D4"] = "Référence"
    reference_suffix = hashlib.sha256(token_id.encode("utf-8")).hexdigest()[:8].upper()
    ws_summary["E4"] = f"FIN-{year}-{reference_suffix}"
    ws_summary["A6"] = "Chiffre d'affaires"
    ws_summary["B6"] = "=Résultat!C3"
    ws_summary["C6"] = "Charges"
    ws_summary["D6"] = "=Résultat!C9"
    ws_summary["E6"] = "Résultat net"
    ws_summary["F6"] = "=Résultat!C10"
    for coordinate in ("A6", "C6", "E6"):
        ws_summary[coordinate].font = Font(bold=True, color="FFFFFF")
        ws_summary[coordinate].fill = PatternFill("solid", fgColor="5B9BD5")
    for coordinate in ("B6", "D6", "F6"):
        ws_summary[coordinate].font = Font(bold=True, size=12)
        ws_summary[coordinate].number_format = _CURRENCY_FORMAT
        ws_summary[coordinate].fill = PatternFill("solid", fgColor="DDEBF7")
    ws_summary.merge_cells("A8:F14")
    ws_summary["A8"] = (
        summary_text or "Synthèse financière préparée pour le comité de direction."
    ).strip()[:3000]
    ws_summary["A8"].alignment = Alignment(wrap_text=True, vertical="top")
    ws_summary["A8"].font = Font(name="Aptos", size=11)
    ws_summary["A16"] = (
        "CONFIDENTIEL — Diffusion strictement limitée au comité de direction. "
        "Toute transmission non autorisée doit être signalée."
    )
    ws_summary["A16"].font = Font(bold=True, italic=True, color="9C0006")
    ws_summary.merge_cells("A16:F16")
    for column in ("A", "B", "C", "D", "E", "F"):
        ws_summary.column_dimensions[column].width = 22

    metadata = {
        "schema_version": "3",
        "workbook_profile": "finance-reporting/3.2",
        "generated_at_utc": datetime.now(timezone.utc).isoformat(),
        "classification": "CONFIDENTIEL",
        "scenario": "financial_accounting",
        "company": company_name,
        "fiscal_year": str(year),
        "journal_lines": str(len(lines)),
        "journal_total_debit": str(
            sum((line.debit for line in lines), Decimal("0.00"))
        ),
        "journal_total_credit": str(
            sum((line.credit for line in lines), Decimal("0.00"))
        ),
    }
    ws_meta.append(["Clé", "Valeur"])
    for key, value in metadata.items():
        ws_meta.append([key, value])
    ws_meta.sheet_state = "veryHidden"

    for worksheet in wb.worksheets:
        worksheet.sheet_view.showGridLines = worksheet.title in {"Journal", "Balance"}
        worksheet.sheet_properties.pageSetUpPr.fitToPage = True
        worksheet.page_setup.fitToWidth = 1
        worksheet.page_setup.fitToHeight = 0

    attach_placeholder_image(wb, ws_summary, token_url)
    return wb
