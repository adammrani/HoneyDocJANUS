from __future__ import annotations

import re
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
SOURCE_MD = ROOT / "JANUS_Cahier_des_charges_v1.0.md"
OUTPUT_DOCX = ROOT / "JANUS_Cahier_des_charges_v1.0.docx"
ARCHITECTURE_PNG = ROOT / "assets" / "janus_architecture_cible.png"

CONTENT_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120
NAVY = "0B2545"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
MUTED = "5D6875"
LIGHT_GRAY = "F2F4F7"
CALLOUT = "F4F6F9"
TEAL = "227C7A"
GOLD = "A86F00"
WHITE = "FFFFFF"
BLACK = "111111"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.first_child_found_in("w:tcW")
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths: list[int]) -> None:
    if sum(widths) != CONTENT_WIDTH_DXA:
        raise ValueError(f"Table width must total {CONTENT_WIDTH_DXA}: {widths}")
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(CONTENT_WIDTH_DXA))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT_DXA))
    tbl_ind.set(qn("w:type"), "dxa")
    layout = tbl_pr.first_child_found_in("w:tblLayout")
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            set_cell_width(cell, widths[index])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def mark_header_row(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def set_repeat_table_header(row) -> None:
    mark_header_row(row)


def font_run(run, *, size=None, color=None, bold=None, italic=None, name="Calibri") -> None:
    run.font.name = name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def shade_paragraph(paragraph, fill: str, left_border: str | None = None) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)
    if left_border:
        p_bdr = OxmlElement("w:pBdr")
        left = OxmlElement("w:left")
        left.set(qn("w:val"), "single")
        left.set(qn("w:sz"), "18")
        left.set(qn("w:space"), "8")
        left.set(qn("w:color"), left_border)
        p_bdr.append(left)
        p_pr.append(p_bdr)


def add_field(paragraph, instruction: str, fallback: str = "1") -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = fallback
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])
    font_run(run, size=9, color=MUTED)


def add_hyperlink(paragraph, text: str, url: str):
    part = paragraph.part
    rel_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.extend([color, underline])
    new_run.append(r_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    new_run.append(text_node)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


INLINE_RE = re.compile(r"(\*\*.+?\*\*|`.+?`|https?://\S+)")


def add_inline(paragraph, text: str, *, size: float | None = None, color: str = BLACK) -> None:
    cursor = 0
    for match in INLINE_RE.finditer(text):
        if match.start() > cursor:
            run = paragraph.add_run(text[cursor : match.start()])
            font_run(run, size=size, color=color)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            font_run(run, size=size, color=color, bold=True)
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            font_run(run, size=(size or 11) - 0.5, color=DARK_BLUE, name="Consolas")
        else:
            clean = token.rstrip(".,;)")
            trailing = token[len(clean) :]
            add_hyperlink(paragraph, clean, clean)
            if trailing:
                run = paragraph.add_run(trailing)
                font_run(run, size=size, color=color)
        cursor = match.end()
    if cursor < len(text):
        run = paragraph.add_run(text[cursor:])
        font_run(run, size=size, color=color)


def create_numbering(document: Document, kind: str) -> int:
    numbering = document.part.numbering_part.element
    abstract_ids = [int(x.get(qn("w:abstractNumId"))) for x in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(x.get(qn("w:numId"))) for x in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids, default=0) + 1
    num_id = max(num_ids, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "bullet" if kind == "bullet" else "decimal")
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), "•" if kind == "bullet" else "%1.")
    lvl_jc = OxmlElement("w:lvlJc")
    lvl_jc.set(qn("w:val"), "left")
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "720")
    tabs.append(tab)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "720")
    ind.set(qn("w:hanging"), "360")
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:after"), "160")
    spacing.set(qn("w:line"), "280")
    spacing.set(qn("w:lineRule"), "auto")
    p_pr.extend([tabs, ind, spacing])
    lvl.extend([start, num_fmt, lvl_text, lvl_jc, p_pr])
    abstract.append(lvl)
    first_num = numbering.find(qn("w:num"))
    if first_num is None:
        numbering.append(abstract)
    else:
        first_num.addprevious(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def apply_num(paragraph, num_id: int) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, num])
    p_pr.append(num_pr)


def set_style(style, *, font_size, color, bold, before, after, line=1.0, keep=False) -> None:
    style.font.name = "Calibri"
    style._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), "Calibri")
    style._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), "Calibri")
    style.font.size = Pt(font_size)
    style.font.color.rgb = RGBColor.from_string(color)
    style.font.bold = bold
    pf = style.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    if keep:
        pf.keep_with_next = True


def configure_document(document: Document) -> tuple[int, int]:
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = document.styles
    set_style(styles["Normal"], font_size=11, color=BLACK, bold=False, before=0, after=6, line=1.10)
    set_style(styles["Heading 1"], font_size=16, color=BLUE, bold=True, before=16, after=8, line=1.0, keep=True)
    set_style(styles["Heading 2"], font_size=13, color=BLUE, bold=True, before=12, after=6, line=1.0, keep=True)
    set_style(styles["Heading 3"], font_size=12, color=DARK_BLUE, bold=True, before=8, after=4, line=1.0, keep=True)
    styles["Heading 1"].paragraph_format.page_break_before = False

    if "Table Citation" not in styles:
        citation = styles.add_style("Table Citation", WD_STYLE_TYPE.PARAGRAPH)
    else:
        citation = styles["Table Citation"]
    set_style(citation, font_size=9, color=MUTED, bold=False, before=4, after=4, line=1.0)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.paragraph_format.space_after = Pt(0)
    run = header.add_run("JANUS  ·  Cahier des charges de redémarrage")
    font_run(run, size=8.5, color=MUTED, bold=True)

    footer = section.footer.paragraphs[0]
    footer.paragraph_format.tab_stops.add_tab_stop(Inches(6.5), WD_TAB_ALIGNMENT.RIGHT)
    footer.paragraph_format.space_before = Pt(0)
    footer.paragraph_format.space_after = Pt(0)
    left = footer.add_run("JANUS  ·  Version 1.0")
    font_run(left, size=8.5, color=MUTED)
    tab = footer.add_run("\tPage ")
    font_run(tab, size=8.5, color=MUTED)
    add_field(footer, "PAGE")
    slash = footer.add_run(" / ")
    font_run(slash, size=8.5, color=MUTED)
    add_field(footer, "NUMPAGES")

    settings = document.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")
    return create_numbering(document, "bullet"), create_numbering(document, "decimal")


def draw_wrapped(draw, box, text, font, color, align="center", spacing=6):
    x1, y1, x2, y2 = box
    words = text.split()
    lines, line = [], ""
    max_width = x2 - x1 - 34
    for word in words:
        trial = f"{line} {word}".strip()
        if draw.textbbox((0, 0), trial, font=font)[2] <= max_width:
            line = trial
        else:
            if line:
                lines.append(line)
            line = word
    if line:
        lines.append(line)
    line_height = draw.textbbox((0, 0), "Ag", font=font)[3] + spacing
    total = len(lines) * line_height
    y = y1 + ((y2 - y1) - total) / 2
    for item in lines:
        width = draw.textbbox((0, 0), item, font=font)[2]
        x = x1 + ((x2 - x1) - width) / 2 if align == "center" else x1 + 18
        draw.text((x, y), item, font=font, fill=color)
        y += line_height


def arrow(draw, start, end, color="#5D6875", width=5):
    draw.line([start, end], fill=color, width=width)
    x, y = end
    if abs(end[0] - start[0]) >= abs(end[1] - start[1]):
        points = [(x, y), (x - 14, y - 9), (x - 14, y + 9)]
    else:
        points = [(x, y), (x - 9, y - 14), (x + 9, y - 14)]
    draw.polygon(points, fill=color)


def poly_arrow(draw, points, color="#5D6875", width=5):
    draw.line(points, fill=color, width=width, joint="curve")
    x, y = points[-1]
    px, py = points[-2]
    if abs(x - px) >= abs(y - py):
        direction = 1 if x > px else -1
        head = [(x, y), (x - 14 * direction, y - 9), (x - 14 * direction, y + 9)]
    else:
        direction = 1 if y > py else -1
        head = [(x, y), (x - 9, y - 14 * direction), (x + 9, y - 14 * direction)]
    draw.polygon(head, fill=color)


def make_architecture_diagram(path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    width, height = 2000, 1120
    image = Image.new("RGB", (width, height), "#FFFFFF")
    draw = ImageDraw.Draw(image)
    font_regular = ImageFont.truetype(r"C:\Windows\Fonts\calibri.ttf", 31)
    font_bold = ImageFont.truetype(r"C:\Windows\Fonts\calibrib.ttf", 32)
    font_small = ImageFont.truetype(r"C:\Windows\Fonts\calibri.ttf", 25)

    draw.rounded_rectangle((30, 30, 1970, 335), radius=24, fill="#EEF4FA", outline="#C6D7E8", width=3)
    draw.text((60, 48), "PLAN DE CONTRÔLE", font=font_bold, fill="#0B2545")
    control = [
        ((80, 125, 330, 265), "Opérateur / CLI", NAVY),
        ((405, 125, 705, 265), "Campagne et registre", BLUE),
        ((780, 125, 1110, 265), "Contenu synthétique", TEAL),
        ((1185, 125, 1515, 265), "Token + artefact", GOLD),
        ((1590, 125, 1915, 265), "Validation + déploiement", NAVY),
    ]
    for box, label, fill in control:
        draw.rounded_rectangle(box, radius=18, fill=f"#{fill}", outline=f"#{fill}")
        draw_wrapped(draw, box, label, font_bold, "#FFFFFF")
    for left, right in zip(control, control[1:]):
        arrow(draw, (left[0][2] + 10, 195), (right[0][0] - 10, 195))

    draw.rounded_rectangle((30, 375, 1970, 650), radius=24, fill="#F8F5EE", outline="#E7DCC1", width=3)
    draw.text((60, 393), "SOURCES DE PREUVE", font=font_bold, fill="#0B2545")
    sources = [
        ((80, 445, 420, 535), "Windows Security / Sysmon", BLUE),
        ((510, 545, 850, 635), "Entrée SMB / RDP / VPN", TEAL),
        ((960, 485, 1285, 595), "Wazuh Manager / Indexer", NAVY),
        ((1410, 485, 1915, 595), "Canarytokens : webhook + history", GOLD),
    ]
    for box, label, fill in sources:
        draw.rounded_rectangle(box, radius=18, fill=f"#{fill}", outline=f"#{fill}")
        draw_wrapped(draw, box, label, font_regular, "#FFFFFF")
    poly_arrow(draw, [(420, 490), (900, 490), (940, 535)])
    poly_arrow(draw, [(850, 590), (900, 590), (940, 555)])

    draw.rounded_rectangle((30, 690, 1970, 1090), radius=24, fill="#F3F7F6", outline="#C9DEDA", width=3)
    draw.text((60, 708), "PLAN DE PREUVE — RAW FIRST", font=font_bold, fill="#0B2545")
    ingest = ((800, 750, 1200, 845), "Ingestion brute", NAVY)
    evidence = [
        ((80, 900, 410, 1025), "Raw store + SHA-256", TEAL),
        ((490, 900, 820, 1025), "Normalisation versionnée", BLUE),
        ((900, 900, 1240, 1025), "Corrélation déterministe", GOLD),
        ((1320, 900, 1915, 1025), "Timeline + Evidence Bundle", NAVY),
    ]
    draw.rounded_rectangle(ingest[0], radius=18, fill=f"#{ingest[2]}", outline=f"#{ingest[2]}")
    draw_wrapped(draw, ingest[0], ingest[1], font_regular, "#FFFFFF")
    for box, label, fill in evidence:
        draw.rounded_rectangle(box, radius=18, fill=f"#{fill}", outline=f"#{fill}")
        draw_wrapped(draw, box, label, font_regular, "#FFFFFF")
    for left, right in zip(evidence, evidence[1:]):
        arrow(draw, (left[0][2] + 10, 962), (right[0][0] - 10, 962))
    poly_arrow(draw, [(1120, 595), (1120, 720), (1000, 740)])
    poly_arrow(draw, [(1660, 595), (1660, 700), (1200, 700), (1200, 795)])
    poly_arrow(draw, [(800, 798), (245, 798), (245, 880)])
    draw.text((90, 1045), "Secrets hors registre · erreurs de parsing conservées et rejouables", font=font_small, fill="#5D6875")
    image.save(path, format="PNG", optimize=True)


def column_widths(headers: list[str]) -> list[int]:
    count = len(headers)
    if count == 2:
        first = 2700 if headers[0] in {"Élément de `legacy`", "Élément historique", "Donnée", "Entité", "Risque"} else 2300
        return [first, CONTENT_WIDTH_DXA - first]
    if count == 3:
        return [1450, 1900, 6010]
    if count == 4:
        return [1350, 1850, 2850, 3310]
    return [CONTENT_WIDTH_DXA // count] * (count - 1) + [CONTENT_WIDTH_DXA - (CONTENT_WIDTH_DXA // count) * (count - 1)]


def clean_cell(text: str) -> str:
    return text.replace("**", "").replace("`", "").strip()


def set_row_cant_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:cantSplit")) is None:
        tr_pr.append(OxmlElement("w:cantSplit"))


def add_markdown_table(document: Document, rows: list[list[str]]) -> None:
    headers = [clean_cell(x) for x in rows[0]]
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = 0
    hdr = table.rows[0]
    set_row_cant_split(hdr)
    for index, text in enumerate(headers):
        cell = hdr.cells[index]
        set_cell_shading(cell, LIGHT_GRAY)
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.keep_with_next = True
        add_inline(p, text, size=9.5, color=NAVY)
        for run in p.runs:
            run.bold = True
    for raw_row in rows[1:]:
        row = table.add_row()
        set_row_cant_split(row)
        for index, text in enumerate(raw_row):
            p = row.cells[index].paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.05
            add_inline(p, clean_cell(text), size=9.2)
    set_repeat_table_header(hdr)
    set_table_geometry(table, column_widths(headers))
    after = document.add_paragraph()
    after.paragraph_format.space_after = Pt(2)


def add_cover(document: Document) -> None:
    document.add_paragraph().paragraph_format.space_after = Pt(34)
    kicker = document.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.LEFT
    kicker.paragraph_format.space_after = Pt(8)
    run = kicker.add_run("CAHIER DES CHARGES  ·  ARCHITECTURE CIBLE")
    font_run(run, size=10.5, color=GOLD, bold=True)

    title = document.add_paragraph()
    title.paragraph_format.space_after = Pt(8)
    run = title.add_run("JANUS")
    font_run(run, size=34, color=NAVY, bold=True)

    subtitle = document.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(24)
    run = subtitle.add_run("Redémarrage fondé sur la preuve\nHoney-Documents Dynamiques assistés par l’IA")
    font_run(run, size=17, color=DARK_BLUE, bold=False)

    metadata = [
        ("Version", "1.0"),
        ("Date", "4 août 2026"),
        ("Statut", "Proposition à valider avant toute implémentation"),
        ("Base", "Archive JANUS-HANDOFF-2026-08-04.zip, revue prioritaire de legacy"),
    ]
    for label, value in metadata:
        p = document.add_paragraph()
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(f"{label} : ")
        font_run(r, size=10.5, color=NAVY, bold=True)
        r = p.add_run(value)
        font_run(r, size=10.5, color=BLACK)

    document.add_paragraph().paragraph_format.space_after = Pt(18)
    lead = document.add_paragraph()
    lead.paragraph_format.left_indent = Inches(0.18)
    lead.paragraph_format.right_indent = Inches(0.18)
    lead.paragraph_format.space_before = Pt(8)
    lead.paragraph_format.space_after = Pt(8)
    shade_paragraph(lead, CALLOUT, BLUE)
    run = lead.add_run(
        "Décision proposée — JANUS devient un système de preuve de cyberdéception. "
        "Le premier jalon démontre un seul DOCX sur deux machines Windows, avec Wazuh et un Canarytoken officiel."
    )
    font_run(run, size=11.5, color=NAVY, bold=True)

    document.add_paragraph().paragraph_format.space_after = Pt(34)
    note = document.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = note.add_run("Document de conception — aucun code métier ne doit précéder sa validation.")
    font_run(run, size=9.5, color=MUTED, italic=True)
    note.add_run().add_break(WD_BREAK.PAGE)


def add_summary(document: Document, headings: list[str], decimal_id: int) -> None:
    p = document.add_paragraph("Sommaire", style="Heading 1")
    p.paragraph_format.space_before = Pt(0)
    intro = document.add_paragraph(
        "Les sections ci-dessous forment la référence de cadrage du nouveau projet."
    )
    intro.paragraph_format.space_after = Pt(10)
    for heading in headings:
        p = document.add_paragraph()
        apply_num(p, decimal_id)
        title = re.sub(r"^\d+\.\s*", "", heading)
        add_inline(p, title)
    document.add_page_break()


def build_document() -> None:
    make_architecture_diagram(ARCHITECTURE_PNG)
    source = SOURCE_MD.read_text(encoding="utf-8").splitlines()
    document = Document()
    bullet_id, decimal_id = configure_document(document)
    document.core_properties.title = "JANUS — Cahier des charges de redémarrage et architecture cible"
    document.core_properties.subject = "Architecture et exigences fondées sur la preuve"
    document.core_properties.author = "Projet JANUS"
    document.core_properties.keywords = "JANUS, cyberdéception, Wazuh, Canarytokens, preuve"
    add_cover(document)

    top_headings = [line[3:].strip() for line in source if line.startswith("## ")]
    add_summary(document, top_headings, decimal_id)

    start = next(i for i, line in enumerate(source) if line.startswith("## 1."))
    index = start
    table_buffer: list[list[str]] = []
    last_list_kind: str | None = None
    active_decimal_id: int | None = None
    active_bullet_id: int | None = None

    def flush_table():
        nonlocal table_buffer
        if table_buffer:
            rows = [row for row in table_buffer if not all(re.fullmatch(r":?-{3,}:?", c.strip()) for c in row)]
            add_markdown_table(document, rows)
            table_buffer = []

    while index < len(source):
        line = source[index].rstrip()
        if line.startswith("|") and line.endswith("|"):
            last_list_kind = None
            cells = [c.strip() for c in line.strip("|").split("|")]
            table_buffer.append(cells)
            index += 1
            continue
        flush_table()

        if not line or line == "---":
            last_list_kind = None
            index += 1
            continue
        if line.startswith("## "):
            last_list_kind = None
            heading_text = line[3:].strip()
            if heading_text.startswith(("7.", "12.")):
                document.add_page_break()
            document.add_paragraph(heading_text, style="Heading 1")
        elif line.startswith("### "):
            last_list_kind = None
            document.add_paragraph(line[4:].strip(), style="Heading 2")
        elif line.startswith("#### "):
            last_list_kind = None
            document.add_paragraph(line[5:].strip(), style="Heading 3")
        elif re.match(r"^\d+\.\s+", line):
            if last_list_kind != "ordered":
                active_decimal_id = create_numbering(document, "decimal")
            p = document.add_paragraph()
            apply_num(p, active_decimal_id)
            add_inline(p, re.sub(r"^\d+\.\s+", "", line))
            last_list_kind = "ordered"
        elif line.startswith("- "):
            if last_list_kind != "bullet":
                active_bullet_id = create_numbering(document, "bullet")
            p = document.add_paragraph()
            apply_num(p, active_bullet_id)
            add_inline(p, line[2:])
            last_list_kind = "bullet"
        elif line.startswith("![Architecture"):
            last_list_kind = None
            p = document.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run()
            picture = run.add_picture(str(ARCHITECTURE_PNG), width=Inches(6.45))
            doc_pr = picture._inline.docPr
            doc_pr.set("descr", "Architecture cible JANUS séparant plan de contrôle, sources et plan de preuve raw-first")
            caption = document.add_paragraph(style="Table Citation")
            caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = caption.add_run("Figure 1 — Architecture cible JANUS : séparation contrôle, sources et preuve raw-first.")
            font_run(r, size=9, color=MUTED, italic=True)
        elif line.startswith("> "):
            last_list_kind = None
            p = document.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.18)
            p.paragraph_format.right_indent = Inches(0.18)
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(10)
            shade_paragraph(p, CALLOUT, BLUE)
            add_inline(p, line[2:], color=NAVY)
        else:
            last_list_kind = None
            p = document.add_paragraph()
            add_inline(p, line)
        index += 1
    flush_table()

    OUTPUT_DOCX.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT_DOCX)


if __name__ == "__main__":
    build_document()
