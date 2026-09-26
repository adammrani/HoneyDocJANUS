from collections import defaultdict
from decimal import Decimal

from openpyxl import load_workbook

from src.janus.generators.accounting_workbook import build_accounting_workbook
from src.janus.generators.xlsx_canary import inspect_external_image_beacon


def test_accounting_workbook_is_balanced_and_macro_free(tmp_path):
    output = tmp_path / "budget_confidentiel.xlsx"
    workbook = build_accounting_workbook(
        summary_text="Synthèse financière confidentielle générée pour le comité.",
        token_id="tok_janus_accounting_001",
        token_url="http://localhost:8000/ping/tok_janus_accounting_001",
        company_name="JANUS Industrie SA",
        fiscal_year=2026,
    )
    workbook.save(output)

    beacon = inspect_external_image_beacon(output)
    assert beacon["valid"] is True
    assert beacon["external_image_targets"] == [
        "http://localhost:8000/ping/tok_janus_accounting_001"
    ]
    assert beacon["macro_parts"] == []

    # openpyxl ne sait pas relire une image dont la relation est externe.
    # On valide donc le contenu comptable sur une sauvegarde sans beacon,
    # tandis que le ZIP tokenisé est contrôlé ci-dessus.
    readable_output = tmp_path / "budget_confidentiel_structure.xlsx"
    readable = build_accounting_workbook(
        summary_text="Synthèse financière confidentielle générée pour le comité.",
        token_id="tok_janus_accounting_001",
        token_url="",
        company_name="JANUS Industrie SA",
        fiscal_year=2026,
    )
    readable.save(readable_output)
    loaded = load_workbook(readable_output, data_only=False, keep_vba=False)
    assert loaded.sheetnames == [
        "Synthèse",
        "Journal",
        "Balance",
        "Résultat",
        "Bilan",
        "Budget mensuel",
        "_Document_Metadata",
    ]
    assert loaded.vba_archive is None
    assert loaded._external_links == []
    assert loaded["_Document_Metadata"].sheet_state == "veryHidden"
    metadata_values = {
        row[0].value: row[1].value
        for row in loaded["_Document_Metadata"].iter_rows(min_row=2)
    }
    assert "token_id" not in metadata_values
    assert "token_url" not in metadata_values
    assert metadata_values["workbook_profile"] == "finance-reporting/3.2"
    assert loaded["Journal"].auto_filter.ref is None
    assert len(loaded["Journal"].tables) == 1

    totals = defaultdict(lambda: [Decimal("0.00"), Decimal("0.00")])
    journal = loaded["Journal"]
    for row in journal.iter_rows(min_row=2, values_only=True):
        document_no = row[1]
        totals[document_no][0] += Decimal(str(row[5] or 0))
        totals[document_no][1] += Decimal(str(row[6] or 0))

    assert totals
    assert all(debit == credit for debit, credit in totals.values())
    assert sum(value[0] for value in totals.values()) == sum(
        value[1] for value in totals.values()
    )
    assert loaded["Balance"]["C2"].value.startswith("=SUMIF(")
    assert loaded["Bilan"]["C12"].value.startswith("=")


def test_same_token_produces_same_financial_values():
    first = build_accounting_workbook("Résumé A", "stable-token", fiscal_year=2026)
    second = build_accounting_workbook("Résumé B", "stable-token", fiscal_year=2026)

    first_values = [
        row for row in first["Journal"].iter_rows(min_row=2, values_only=True)
    ]
    second_values = [
        row for row in second["Journal"].iter_rows(min_row=2, values_only=True)
    ]
    assert first_values == second_values
